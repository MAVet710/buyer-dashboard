from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from modules.coman.models import utc_now
from tests.test_operator_acceptance_retail_extraction import _seeded_client


def _product_metrics(payload: dict, product_id: str) -> dict[str, float]:
    row = next((item for item in payload.get("products", []) if item.get("product_id") == product_id), None)
    return {
        "quantity": float(row.get("quantity", 0)) if row else 0.0,
        "net_sales": float(row.get("net_sales", 0)) if row else 0.0,
    }


def test_retail_trends_and_audit_exports_match_the_exact_source_transactions():
    client, headers = _seeded_client(role="operator", user_id="functional-output-reporter")
    try:
        inventory = client.get("/api/v1/inventory/retail/packages", headers=headers)
        assert inventory.status_code == 200, inventory.text
        source = next(row for row in inventory.json()["items"] if row["available"] > 0)
        product_id = source["product_id"]

        before = client.get("/api/v1/retail-insights/trends?days=30", headers=headers)
        assert before.status_code == 200, before.text
        before_payload = before.json()
        before_product = _product_metrics(before_payload, product_id)
        before_units = float(before_payload["summary"]["units"])
        before_sales = float(before_payload["summary"]["net_sales"])

        sale_quantity = 3.0
        sale_net = 75.0
        batch = f"FO-SALES-{uuid4().hex[:8]}"
        imported = client.post(
            "/api/v1/inventory/retail/sales/import",
            headers=headers,
            json={
                "source_system": "FunctionalOutputPOS",
                "import_batch_id": batch,
                "lines": [{
                    "source_record_id": f"{batch}-LINE-1",
                    "sold_at": utc_now().isoformat(),
                    "quantity": sale_quantity,
                    "product_id": product_id,
                    "sku": source["sku"],
                    "product_name": source["product_name"],
                    "net_sales": sale_net,
                }],
            },
        )
        assert imported.status_code == 201, imported.text
        assert imported.json()["imported"] == 1

        after = client.get("/api/v1/retail-insights/trends?days=30", headers=headers)
        assert after.status_code == 200, after.text
        after_payload = after.json()
        after_product = _product_metrics(after_payload, product_id)
        assert float(after_payload["summary"]["units"]) == before_units + sale_quantity
        assert float(after_payload["summary"]["net_sales"]) == before_sales + sale_net
        assert after_product["quantity"] == before_product["quantity"] + sale_quantity
        assert after_product["net_sales"] == before_product["net_sales"] + sale_net
        assert float(after_payload["summary"]["average_daily_sales"]) == (before_sales + sale_net) / 30

        package_id = "1A4000000000000000098001"
        lot_code = "FO-AUDIT-LOT-0001"
        receipt = client.post(
            "/api/v1/inventory/retail/receipts",
            headers=headers,
            json={
                "product_id": product_id,
                "lot_code": lot_code,
                "package_id": package_id,
                "quantity": 7,
                "unit": source["unit"],
                "location": "FO-VAULT",
                "source_name": "Functional Output Vendor",
                "manifest_reference": "FO-MANIFEST-0001",
                "lab_testing_state": "TestPassed",
                "coa_reference": "FO-COA-0001",
                "notes": "Functional output audit export fixture",
            },
        )
        assert receipt.status_code == 201, receipt.text
        lot_id = receipt.json()["lot_id"]

        audit = client.post(
            "/api/v1/inventory/retail/audits",
            headers=headers,
            json={
                "audit_number": "FO-AUDIT-0001",
                "scope_label": "Functional output exact export",
                "notes": "Validate workbook values, not only download success",
                "lot_ids": [lot_id],
                "blind_count": False,
                "recount_tolerance": 0,
            },
        )
        assert audit.status_code == 201, audit.text
        audit_id = audit.json()["id"]
        detail = client.get(f"/api/v1/inventory/retail/audits/{audit_id}", headers=headers)
        assert detail.status_code == 200, detail.text
        line = detail.json()["lines"][0]
        assert line["expected_quantity"] == 7

        counted = client.post(
            f"/api/v1/inventory/retail/audits/{audit_id}/counts",
            headers=headers,
            json={"counts": [{
                "line_id": line["id"],
                "counted_quantity": 6,
                "reason": "Physical count variance",
                "notes": "One unit missing in functional-output fixture",
            }]},
        )
        assert counted.status_code == 200, counted.text
        counted_line = counted.json()["lines"][0]
        assert counted_line["expected_quantity"] == 7
        assert counted_line["counted_quantity"] == 6
        assert counted_line["variance_quantity"] == -1

        csv_report = client.get(f"/api/v1/inventory/retail/audits/{audit_id}/report.csv", headers=headers)
        assert csv_report.status_code == 200, csv_report.text
        assert lot_code in csv_report.text
        assert package_id in csv_report.text
        assert "Physical count variance" in csv_report.text

        xlsx_report = client.get(f"/api/v1/inventory/retail/audits/{audit_id}/report.xlsx", headers=headers)
        assert xlsx_report.status_code == 200, xlsx_report.text
        workbook = load_workbook(BytesIO(xlsx_report.content), data_only=True)
        assert workbook.sheetnames == ["Summary", "Audit Detail", "Activity"]

        summary_sheet = workbook["Summary"]
        summary_headers = [cell.value for cell in summary_sheet[1]]
        summary_values = [cell.value for cell in summary_sheet[2]]
        summary_row = dict(zip(summary_headers, summary_values))
        assert summary_row["Audit"] == "FO-AUDIT-0001"
        assert summary_row["Status"] == "in_progress"
        assert summary_row["Scope"] == "Functional output exact export"
        assert summary_row["Created By"] == "functional-output-reporter"

        detail_sheet = workbook["Audit Detail"]
        headers_row = [cell.value for cell in detail_sheet[1]]
        values_row = [cell.value for cell in detail_sheet[2]]
        exported = dict(zip(headers_row, values_row))
        assert exported["Lot / Batch"] == lot_code
        assert exported["METRC Package"] == package_id
        assert exported["Location"] == "FO-VAULT"
        assert exported["Expected"] == 7
        assert exported["Final Count"] == 6
        assert exported["Variance"] == -1
        assert exported["Reason"] == "Physical count variance"
        assert exported["Notes"] == "One unit missing in functional-output fixture"
        assert exported["Counted By"] == "functional-output-reporter"
        assert exported["Cost Impact"] == -float(exported["Unit Cost"] or 0)
        assert exported["Revenue Impact"] == -float(exported["Retail Price"] or 0)
    finally:
        client.close()
        from backend.app.main import app
        app.dependency_overrides.clear()
